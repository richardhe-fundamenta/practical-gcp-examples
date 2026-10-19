# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Richard He
import io
import json

import httpx
import jwt
import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from starlette.testclient import TestClient

from sharepoint_mcp_bridge import audit, graph, office
from sharepoint_mcp_bridge.server import app, mcp

READ_TOOLS = {"search_documents", "semantic_search", "list_sites_and_drives", "list_folder", "list_recent_and_shared",
              "read_document", "get_file_details", "compare_versions", "get_permissions", "query_list"}


def entra_token(**claims) -> str:
    return jwt.encode({"oid": "oid-1", "tid": "tid-1", "upn": "adele.vance@contoso.com", "scp": "Files.ReadWrite.All",
                       **claims}, "k" * 32, algorithm="HS256")


def docx_with_runs() -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Limit: ")
    bold = p.add_run("EUR 50,000,000")
    bold.bold = True
    p.add_run(" pending.")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Owner", "Status"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_edit_preserves_run_formatting_and_flexible_match():
    data, stats = office.edit_docx(
        docx_with_runs(),
        office.pairs([{"find": "_EUR_50_000_000_", "replace": "EUR 75,000,000"}], "pending", "approved"),
        append_heading="Sign-off", append_paragraphs=["Done."], table_rows=[["Ana", "Open"]])
    assert stats["not_found"] == [] and stats["table_rows_appended"] == 1
    runs = Document(io.BytesIO(data)).paragraphs[0].runs
    assert [r.text for r in runs] == ["Limit: ", "EUR 75,000,000", " approved."]
    assert runs[1].bold and not runs[0].bold
    text = office.extract_text("a.docx", data)
    assert "## Sign-off" in text and "Ana | Open" in text


def test_replacement_across_runs_and_flex_regex():
    rx = office.flex_regex("_Pending_Sign_Off_")
    assert rx.search("Status: pending sign-off today")
    assert not rx.search("Status: pendingsign off")
    assert office.pairs({"a": "b"}) == [("a", "b")]


def test_xlsx_one_pass_and_formula_fallback():
    data, stats = office.edit_xlsx(None, sheet="Close", initial_rows=[["Item", "Plan", "Actual"], ["Rent", 10, 12]],
                                   cell_updates={"D1": "Variance", "D2": "=C2-B2"}, append_rows=[["Fees", 5, 4]],
                                   add_column={"header": "Owner", "values": ["Ana", "Bo"]},
                                   create_table={"name": "Close Table"})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Close"] and wb["Close"]["E3"].value == "Bo"
    assert "Close_Table" in wb["Close"].tables and stats["table_created"]["range"] == "A1:E3"
    assert "Rent | 10 | 12 | =C2-B2 | Ana" in office.extract_text("t.xlsx", data)
    with pytest.raises(ValueError):
        office.edit_xlsx(data, cell_updates={"1A": 3})


def test_pptx_build_edit_extract():
    data = office.build_file("deck.pptx", title="Kickoff", slides=[{"title": "Agenda", "bullets": ["Scope", "Budget"]}])
    data, stats = office.edit_pptx(data, office.pairs({"Budget": "Costs"}), [{"title": "Next", "bullets": ["Ship"]}])
    assert stats["total_slides"] == 3 and stats["replacements"] == {"Budget": 1}
    text = office.extract_text("deck.pptx", data)
    assert "Costs" in text and "## Slide 3" in text and "Ship" in text


def test_template_fill_reports_unfilled():
    doc = Document()
    doc.add_paragraph("Dear {{ client_name }}, amount {{Amount}}.")
    p = doc.add_paragraph()
    p.add_run("Due {{due")
    p.add_run("_date}} ok")
    buf = io.BytesIO()
    doc.save(buf)
    out, filled, unfilled = office.fill_template("t.docx", buf.getvalue(), {"Client Name": "Contoso", "_due_date_": "May 1"})
    text = office.extract_text("t.docx", out)
    assert "Dear Contoso, amount {{Amount}}." in text and "Due May 1 ok" in text
    assert unfilled == ["Amount"] and len(filled) == 2


def test_compare_two_docx():
    a = office.extract_text("a.docx", office.build_file("a.docx", "Line one\nLine two"))
    b = office.extract_text("b.docx", office.build_file("b.docx", "Line one\nLine 2\nLine three"))
    d = office.diff(a, b, "a", "b")
    assert d["added_lines"] == 2 and d["removed_lines"] == 1 and "+Line three" in d["diff"]


@pytest.mark.parametrize("bad", ["../x.docx", "a/b.docx", "..", "x\\y", "a:b", ""])
def test_sanitizers_reject_traversal(bad):
    with pytest.raises(ValueError):
        office.sanitize_name(bad)
    with pytest.raises(ValueError):
        office.sanitize_folder("Reports/../../etc")
    with pytest.raises(ValueError):
        graph.validate_id("abc/../me")
    with pytest.raises(ValueError):
        graph.site_path("evil.com:/sites/x")
    assert office.sanitize_folder("/Reports/2026/") == "Reports/2026"
    assert graph.drive_path("sharepoint") == "/sites/root/drive"


def test_token_ignores_google_id_token(monkeypatch):
    monkeypatch.delenv("MS_GRAPH_ACCESS_TOKEN", raising=False)
    headers = {}
    monkeypatch.setattr(audit, "get_http_headers", lambda include=None, **_: headers)
    headers["authorization"] = "Bearer " + jwt.encode({"iss": "https://accounts.google.com"}, "k" * 32)
    with pytest.raises(PermissionError):
        audit.request_token()
    monkeypatch.setenv("MS_GRAPH_ACCESS_TOKEN", "local")
    assert audit.request_token() == "local"
    headers["authorization"] = "Bearer " + entra_token()
    assert audit.request_token() == entra_token()


def test_audit_line_masks_upn_and_omits_token(monkeypatch, capsys):
    token = entra_token()
    monkeypatch.setattr(audit, "get_http_headers", lambda include=None, **_: {
        "authorization": f"Bearer {token}", "x-cloud-trace-context": "abc123/1;o=1"})
    monkeypatch.setenv("GOOGLE_CLOUD_PROJECT", "demo")

    @audit.audited("READ")
    async def tool(drive="onedrive"):
        audit.note(item_id="I1")
        return {"ok": 1}

    import asyncio
    assert asyncio.run(tool(drive="sharepoint"))["status"] == "success"
    out = capsys.readouterr().out
    rec = json.loads(out)
    assert token not in out and rec["user"]["upn_masked"] == "ad***@contoso.com"
    assert rec["target"] == {"drive": "sharepoint", "item_id": "I1"} and len(rec["user"]["user_hash"]) == 16
    assert rec["logging.googleapis.com/trace"] == "projects/demo/traces/abc123"


async def test_read_only_annotations_match_read_tools():
    tools = await mcp.list_tools()
    assert len(tools) == 20 and all(t.description for t in tools)
    assert {t.name for t in tools if t.annotations.read_only_hint} == READ_TOOLS


async def test_anonymous_link_rejected_by_default(monkeypatch):
    monkeypatch.setenv("MS_GRAPH_ACCESS_TOKEN", "t")
    monkeypatch.delenv("ALLOW_ANONYMOUS_LINKS", raising=False)
    res = await mcp.call_tool("create_sharing_link", {"file_name": "a.docx", "scope": "anonymous"})
    assert res.structured_content["error_code"] == "ANONYMOUS_LINKS_DISABLED"


async def test_missing_token_is_audited_error(monkeypatch):
    monkeypatch.delenv("MS_GRAPH_ACCESS_TOKEN", raising=False)
    res = await mcp.call_tool("list_folder", {})
    assert res.structured_content["error_code"] == "AUTH_TOKEN_MISSING"


async def test_locked_file_falls_back_to_renamed_copy(monkeypatch):
    calls = []

    def handler(req: httpx.Request):
        calls.append((req.method, req.url.path, req.url.params.get("@microsoft.graph.conflictBehavior")))
        if req.url.path.endswith("/items/F1/content"):
            return httpx.Response(423)
        return httpx.Response(201, json={"id": "F2", "name": "Plan 1.docx", "parentReference": {"driveId": "D"}})

    monkeypatch.setattr(graph, "TRANSPORT", httpx.MockTransport(handler))
    item = {"id": "F1", "name": "Plan.docx", "parentReference": {"driveId": "D", "id": "P"}}
    async with graph.Graph(token="t") as g:
        res = await g.upload_existing(item, b"x")
    assert res["lock_fallback_renamed"] is True and res["file"]["name"] == "Plan 1.docx"
    assert calls[1] == ("PUT", "/v1.0/drives/D/items/P:/Plan.docx:/content", "rename")


def test_healthz_and_security_headers():
    with TestClient(app) as client:
        r = client.get("/healthz")
    assert r.json() == {"status": "ok"}
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["x-frame-options"] == "DENY"
    assert r.headers["cache-control"] == "no-store" and "max-age" in r.headers["strict-transport-security"]


def test_build_csv_xlsx_pptx():
    assert office.build_file("a.csv", rows=[["a", 1]]) == b"a,1\r\n"
    assert load_workbook(io.BytesIO(office.build_file("a.xlsx", "x,y\n1,2"))).active["B2"].value == "2"
    assert Presentation(io.BytesIO(office.build_file("p.pptx", "hello"))).slides
