# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Richard He
"""Generate synthetic demo files for the fictional Halcyon Group.

All people, companies, sites and figures are fictional.
Run: uv run python scripts/generate_demo_files.py
"""

import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Inches
from pptx.util import Pt as PptPt

OUT = Path(__file__).resolve().parent.parent / "demo_assets"
# Same layout as the SharePoint Documents library; Lists/ holds files imported as SharePoint lists.
LEGAL, FINANCE, NETOPS, LISTS = (OUT / d for d in ("Legal", "Finance", "Network Ops", "Lists"))
NAVY = "1F3864"
FIXED_TS = datetime(2026, 10, 1, 9, 0, 0)
FOOTER = "Halcyon Group - fictional demo document - not real data"


# ---------- Word helpers ----------

def new_doc(title, subtitle):
    doc = Document()
    doc.core_properties.author = "Halcyon Group (demo)"
    doc.core_properties.title = title
    doc.core_properties.created = FIXED_TS
    doc.core_properties.modified = FIXED_TS
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string(NAVY)
    section = doc.sections[0]
    section.header.paragraphs[0].text = "HALCYON GROUP"
    section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    section.footer.paragraphs[0].text = FOOTER
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.runs[0].italic = True
    return doc


def para(container, text, style=None, bold=False):
    """Add a paragraph; each {{placeholder}} gets its own run."""
    p = container.add_paragraph(style=style)
    for part in re.split(r"({{\w+}})", text):
        if part:
            p.add_run(part).bold = bold
    return p


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, NAVY)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            para_into_cell(cell, str(text))
    doc.add_paragraph()
    return t


def para_into_cell(cell, text):
    p = cell.paragraphs[0]
    for part in re.split(r"({{\w+}})", text):
        if part:
            p.add_run(part)


# ---------- Legal ----------

def nda_template():
    doc = new_doc("Mutual Non-Disclosure Agreement", "Halcyon Group standard template - Legal")
    para(doc, "This Mutual Non-Disclosure Agreement (the \"Agreement\") is made on {{effective_date}} between:")
    para(doc, "Halcyon Group, 100 Harbour Street, Sydney NSW 2000 (\"Halcyon\"); and", style="List Bullet")
    para(doc, "{{counterparty}}, of {{counterparty_address}} (the \"Counterparty\").", style="List Bullet")
    clauses = [
        ("1. Purpose", "The parties wish to exchange Confidential Information to evaluate a potential business "
         "relationship (the \"Purpose\")."),
        ("2. Confidential Information", "\"Confidential Information\" means any non-public information disclosed "
         "by either party, in any form, that is marked confidential or would reasonably be understood to be "
         "confidential."),
        ("3. Obligations", "Each party shall use the other party's Confidential Information solely for the Purpose, "
         "protect it with at least reasonable care, and disclose it only to personnel who need to know it."),
        ("4. Exclusions", "Obligations do not apply to information that is public, already known to the recipient, "
         "independently developed, or lawfully received from a third party."),
        ("5. Term", "This Agreement remains in force for {{term_years}} years from the Effective Date. Obligations "
         "survive for the same period after termination."),
        ("6. Return of Information", "On request, each party shall return or destroy the other party's "
         "Confidential Information."),
        ("7. Governing Law", "This Agreement is governed by the laws of {{governing_law}}."),
    ]
    for heading, body in clauses:
        doc.add_heading(heading, level=2)
        para(doc, body)
    doc.add_heading("Signatures", level=2)
    table(doc, ["", "Halcyon Group", "{{counterparty}}"], [
        ["Name", "{{halcyon_signatory}}", ""],
        ["Signature", "", ""],
        ["Date", "{{effective_date}}", "{{effective_date}}"],
    ])
    doc.save(LEGAL / "NDA_Template.docx")


def msa(version):
    v4 = version == 4
    doc = new_doc(f"Master Services Agreement - Version {version}",
                  "Between Halcyon Group and Fabrikam Ltd - DRAFT for negotiation")
    table(doc, ["Item", "Detail"], [
        ["Customer", "Halcyon Group"],
        ["Supplier", "Fabrikam Ltd"],
        ["Document version", f"v{version}"],
        ["Status", "Draft - subject to contract"],
    ])
    liability = ("Each party's total aggregate liability under this Agreement is limited to the fees paid or "
                 f"payable in the {'six (6)' if v4 else 'twelve (12)'} months preceding the event giving rise to "
                 "the claim.")
    if v4:
        liability += (" This cap does not apply to liability arising from a Data Breach caused by the Supplier, "
                      "which is uncapped.")
    clauses = [
        ("Definitions", "Capitalised terms have the meanings given in Schedule 1. \"Services\" means the managed "
         "services described in each Statement of Work."),
        ("Services", "The Supplier shall provide the Services in accordance with each Statement of Work, with "
         "reasonable skill and care and in line with good industry practice."),
        ("Fees and Payment", "The Customer shall pay undisputed invoices within "
         f"{'thirty (30)' if v4 else 'forty-five (45)'} days of receipt. Fees are exclusive of applicable taxes."),
        ("Service Levels", "The Supplier shall meet the service levels in Schedule 2. Service credits are the "
         "Customer's sole financial remedy for service level failures."),
        ("Confidentiality", "Each party shall keep the other party's Confidential Information confidential and use "
         "it only to perform this Agreement."),
        ("Intellectual Property", "Each party retains its pre-existing intellectual property. Deliverables created "
         "specifically for the Customer vest in the Customer on payment."),
        ("Limitation of Liability", liability + " Neither party is liable for indirect or consequential loss."),
    ]
    clauses.append(("Term and Termination",
                    "This Agreement starts on the Effective Date and continues for three (3) years. Either party "
                    "may terminate for convenience on "
                    f"{'thirty (30)' if v4 else 'ninety (90)'} days' written notice, or immediately for material "
                    "breach not remedied within 30 days."))
    clauses.append(("General", "This Agreement is governed by the laws of New South Wales, Australia. It is the "
                    "entire agreement between the parties and may be varied only in writing."))
    # Appended last so v3 clause numbers are unchanged in v4.
    if v4:
        clauses.append(("Data Protection and Sub-processors",
                        "The Supplier shall process personal data only on the Customer's documented instructions, "
                        "implement appropriate security measures, notify the Customer of any Data Breach within 48 "
                        "hours, and not engage any sub-processor without the Customer's prior written consent. The "
                        "Supplier remains liable for the acts of its sub-processors."))
    for i, (heading, body) in enumerate(clauses, 1):
        doc.add_heading(f"{i}. {heading}", level=2)
        para(doc, body)
    doc.add_heading("Execution", level=2)
    table(doc, ["", "Halcyon Group", "Fabrikam Ltd"], [
        ["Name", "", ""], ["Title", "", ""], ["Signature", "", ""], ["Date", "", ""]])
    doc.save(LEGAL / f"Fabrikam_MSA_v{version}.docx")


# ---------- Finance ----------

HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF")


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")


def close_tracker():
    wb = Workbook()
    wb.properties.creator = "Halcyon Group (demo)"
    wb.properties.created = FIXED_TS
    ws = wb.active
    ws.title = "Close Tasks"
    header = ["Task ID", "Task", "Owner", "Due Date", "Status", "Notes"]
    rows = [
        ("CL-01", "Bank reconciliations", "Treasury Analyst", date(2026, 11, 3), "Done", "All accounts reconciled"),
        ("CL-02", "AP cut-off and invoice matching", "AP Lead", date(2026, 11, 3), "Done", ""),
        ("CL-03", "AR ageing review", "AR Lead", date(2026, 11, 4), "Done", "No new items over 90 days"),
        ("CL-04", "Accruals review", "Financial Controller", date(2026, 11, 4), "In Progress",
         "Awaiting cloud hosting invoice"),
        ("CL-05", "Fixed asset depreciation run", "Assets Accountant", date(2026, 11, 4), "Done", ""),
        ("CL-06", "Payroll journal", "Payroll Lead", date(2026, 11, 5), "In Progress", "Bonus accrual pending"),
        ("CL-07", "Intercompany reconciliation", "Group Accountant", date(2026, 11, 5), "Not Started",
         "Blocked on Network Ops recharge schedule"),
        ("CL-08", "Prepayments amortisation", "GL Accountant", date(2026, 11, 5), "Done", ""),
        ("CL-09", "Variance analysis", "FP&A Manager", date(2026, 11, 6), "In Progress",
         "Travel and cloud hosting over budget"),
        ("CL-10", "Board report pack", "CFO Office", date(2026, 11, 7), "Not Started", "Needs variance commentary"),
    ]
    ws.append(header)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "yyyy-mm-dd"
    tbl = Table(displayName="CloseTasks", ref=f"A1:F{len(rows) + 1}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    style_header(ws, len(header))
    for col, width in zip("ABCDEF", (10, 34, 22, 12, 14, 42)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    var = wb.create_sheet("Variance")
    var.append(["P&L Line", "Budget", "Actual", "Variance", "Variance %"])
    lines = [
        ("Revenue", 4_200_000, 4_310_000),
        ("Cost of goods sold", 2_310_000, 2_350_000),
        ("Salaries and wages", 980_000, 972_000),
        ("Travel", 45_000, 71_500),
        ("Cloud hosting", 120_000, 158_000),
        ("Software licences", 64_000, 62_500),
        ("Marketing", 85_000, 79_000),
        ("Professional fees", 40_000, 43_000),
        ("Facilities", 72_000, 71_000),
    ]
    for i, (name, budget, actual) in enumerate(lines, 2):
        var.append([name, budget, actual, f"=C{i}-B{i}", f"=IF(B{i}=0,0,D{i}/B{i})"])
    total = len(lines) + 2
    var.append(["Total operating expenses (excl. COGS)", f"=SUM(B4:B{total - 1})", f"=SUM(C4:C{total - 1})",
                f"=C{total}-B{total}", f"=IF(B{total}=0,0,D{total}/B{total})"])
    for row in var.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row[:3]:
            cell.number_format = "#,##0;[Red](#,##0)"
        row[3].number_format = "0.0%"
    for cell in var[total]:
        cell.font = Font(bold=True)
    style_header(var, 5)
    var.column_dimensions["A"].width = 38
    for col in "BCDE":
        var.column_dimensions[col].width = 14
    var.freeze_panes = "A2"
    wb.save(FINANCE / "October_Close_Tracker.xlsx")


def board_report():
    doc = new_doc("Board Report - October 2026", "DRAFT - prepared by the CFO Office for the November board meeting")
    doc.add_heading("1. Executive Summary", level=1)
    para(doc, "October trading was ahead of budget on revenue, with operating expenses slightly above plan. "
         "The month-end close is on track for completion by 7 November.")
    doc.add_heading("2. Financial Summary", level=1)
    table(doc, ["Metric", "Budget", "Actual", "Variance"], [
        ["Revenue", "4,200,000", "4,310,000", "110,000"],
        ["Cost of goods sold", "2,310,000", "2,350,000", "(40,000)"],
        ["Gross margin", "1,890,000", "1,960,000", "70,000"],
        ["Operating expenses", "1,406,000", "1,457,000", "(51,000)"],
    ])
    doc.add_heading("Variance Commentary", level=2)
    para(doc, "[Commentary to be added]")
    doc.add_heading("3. Operational Highlights", level=1)
    for item in ("Network Operations completed 14 planned changes with no customer-impacting incidents.",
                 "Legal finalised the Acme Logistics supply agreement; Fabrikam Ltd MSA v4 under review.",
                 "Core router upgrade programme continues at the Sydney POP in week 43."):
        para(doc, item, style="List Bullet")
    doc.add_heading("4. Risks and Matters for the Board", level=1)
    para(doc, "Cloud hosting costs are trending above budget; a cost optimisation review is proposed for Q4.")
    doc.save(FINANCE / "Board_Report_October.docx")


# ---------- Network operations ----------

CHANGE_HEADER = ["Title", "Summary", "Risk", "ScheduledStart", "Site", "Owner", "Status"]
CHANGES = [
    ("CHG-1042", "Core router OS upgrade - Sydney POP", "High", "2026-10-21", "Sydney POP", "Network Engineering Lead",
     "Pending Approval"),
    ("CHG-1043", "Firewall rule base clean-up - Melbourne DC", "Medium", "2026-10-19", "Melbourne DC",
     "Security Operations", "Approved"),
    ("CHG-1044", "DWDM line card replacement - Brisbane POP", "High", "2026-10-22", "Brisbane POP",
     "Transport Engineer", "Approved"),
    ("CHG-1045", "Edge switch firmware patch - Perth POP", "Low", "2026-10-20", "Perth POP", "Field Technician",
     "Scheduled"),
    ("CHG-1046", "BGP peering migration to new IX port - Sydney IX", "High", "2026-10-23", "Sydney IX",
     "Peering Manager", "Pending Approval"),
    ("CHG-1047", "NTP server replacement - Adelaide DC", "Low", "2026-10-20", "Adelaide DC", "Systems Engineer",
     "Scheduled"),
    ("CHG-1048", "Monitoring agent rollout - all POPs", "Medium", "2026-10-24", "All POPs", "NOC Manager",
     "Approved"),
    ("CHG-1049", "UPS battery replacement - Canberra POP", "Medium", "2026-10-25", "Canberra POP",
     "Facilities Engineer", "Scheduled"),
]


def change_register():
    wb = Workbook()
    wb.properties.creator = "Halcyon Group (demo)"
    wb.properties.created = FIXED_TS
    ws = wb.active
    ws.title = "Change Register"
    ws.append(CHANGE_HEADER)
    for c in CHANGES:
        ws.append([*c[:3], date.fromisoformat(c[3]), *c[4:]])
        ws.cell(row=ws.max_row, column=4).number_format = "yyyy-mm-dd"
    tbl = Table(displayName="ChangeRegister", ref=f"A1:G{len(CHANGES) + 1}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    style_header(ws, len(CHANGE_HEADER))
    for col, width in zip("ABCDEFG", (11, 48, 9, 15, 14, 26, 18)):
        ws.column_dimensions[col].width = width
    wb.save(LISTS / "Change_Register.xlsx")


def method_statement():
    doc = new_doc("Method Statement - CHG-1042", "Core router OS upgrade - Sydney POP - Network Operations")
    table(doc, ["Field", "Value"], [
        ["Change ID", "CHG-1042"], ["Site", "Sydney POP"], ["Risk", "High"],
        ["Window", "2026-10-21 00:00-04:00 AEDT"], ["Owner", "Network Engineering Lead"],
        ["Status", "Pending Approval"],
    ])
    doc.add_heading("1. Scope", level=1)
    para(doc, "Upgrade the operating system on core routers SYD-CR-01 and SYD-CR-02 from release 7.4 to 7.9 to "
         "address vendor security advisories. Routers are upgraded one at a time; traffic is drained to the peer "
         "router before each upgrade.")
    doc.add_heading("2. Pre-checks", level=1)
    for item in ("Confirm a current configuration backup for both routers.",
                 "Stage the configuration diff between 7.4 and 7.9 and have it peer reviewed.",
                 "Run the pre-change validation script and confirm zero errors.",
                 "Record BGP neighbour states, route counts and interface status as a baseline.",
                 "Confirm the NOC is staffed and the vendor support case is open."):
        para(doc, item, style="List Number")
    doc.add_heading("3. Implementation Steps", level=1)
    table(doc, ["Step", "Action", "Owner", "Expected Result", "Duration"], [
        ["1", "Announce start of window to NOC bridge", "Change Manager", "NOC acknowledged", "5 min"],
        ["2", "Drain traffic from SYD-CR-01 (raise IGP metric)", "Network Engineer", "Traffic on SYD-CR-02",
         "15 min"],
        ["3", "Install OS 7.9 on SYD-CR-01 and reload", "Network Engineer", "Router boots on 7.9", "30 min"],
        ["4", "Verify BGP sessions and route counts vs baseline", "Network Engineer", "All sessions Established",
         "15 min"],
        ["5", "Restore traffic to SYD-CR-01", "Network Engineer", "Traffic balanced", "10 min"],
        ["6", "Repeat steps 2-5 for SYD-CR-02", "Network Engineer", "Both routers on 7.9", "70 min"],
        ["7", "Post-change monitoring and close-out", "NOC Manager", "No alarms for 30 min", "30 min"],
    ])
    doc.add_heading("4. Rollback Plan", level=1)
    para(doc, "If any BGP session fails to re-establish within 10 minutes, or route counts differ from the baseline "
         "by more than 2%, keep traffic drained, boot the router from the retained 7.4 image, restore the backed-up "
         "configuration and verify against the baseline. Rollback time per router is approximately 30 minutes.")
    doc.add_heading("5. Communications Plan", level=1)
    table(doc, ["When", "Audience", "Channel", "Message"], [
        ["T-48h", "Service desk, key accounts", "Email", "Planned maintenance notice"],
        ["T-0", "NOC, change team", "Bridge call", "Window start"],
        ["Each step", "NOC", "Bridge call", "Step status"],
        ["T+end", "Service desk, CAB", "Email", "Completion or rollback summary"],
    ])
    doc.add_heading("6. Risk Assessment", level=1)
    table(doc, ["Risk", "Likelihood", "Impact", "Mitigation"], [
        ["Configuration mismatch after upgrade breaks BGP", "Medium", "High",
         "Staged config diff and pre-change validation script"],
        ["Router fails to boot new image", "Low", "High", "Retained 7.4 image; vendor on standby"],
        ["Loss of redundancy during upgrade", "Medium", "Medium", "One router at a time; low-traffic window"],
    ])
    doc.save(NETOPS / "CHG-1042_Method_Statement.docx")


def incident_review():
    doc = new_doc("Post-Incident Review - Core Router Upgrade, Melbourne POP",
                  "Incident INC-2207 - May 2026 - Network Operations")
    table(doc, ["Field", "Value"], [
        ["Incident", "INC-2207"], ["Date", "2026-05-14"], ["Site", "Melbourne POP"],
        ["Related change", "CHG-0918 Core router OS upgrade"], ["Severity", "P2"],
        ["Customer impact", "40 minute BGP outage for transit customers"],
    ])
    doc.add_heading("1. Summary", level=1)
    para(doc, "During a planned core router OS upgrade at the Melbourne POP, BGP sessions to upstream transit "
         "providers failed to re-establish after the router reloaded. Transit traffic was disrupted for 40 minutes "
         "until the configuration was corrected.")
    doc.add_heading("2. Timeline", level=1)
    table(doc, ["Time (AEST)", "Event"], [
        ["01:05", "MEL-CR-01 reloaded on new OS release"],
        ["01:12", "NOC alarms: BGP sessions to transit providers down"],
        ["01:20", "Incident declared, bridge opened"],
        ["01:38", "Root cause identified: route-policy syntax not accepted by new release"],
        ["01:52", "Corrected policy applied, BGP sessions re-established"],
        ["02:30", "Monitoring complete, incident closed"],
    ])
    doc.add_heading("3. Root Cause", level=1)
    para(doc, "A configuration mismatch between the old and new OS releases: the new release rejected a legacy "
         "route-policy statement, leaving BGP neighbours without an outbound policy. The config diff had not been "
         "reviewed before the change, and the rollback trigger was not clearly defined.")
    doc.add_heading("4. Lessons Learned", level=1)
    for item in ("Stage the config diff and run a pre-change validation script before every core router upgrade.",
                 "Define explicit rollback triggers (e.g. BGP not re-established within 10 minutes).",
                 "Upgrade one router at a time with traffic drained to the redundant peer.",
                 "Capture a BGP and route-count baseline before the change for fast comparison."):
        para(doc, item, style="List Bullet")
    doc.add_heading("5. Actions", level=1)
    table(doc, ["Action", "Owner", "Due", "Status"], [
        ["Add pre-change validation script to router upgrade runbook", "Network Engineering Lead", "2026-06-30",
         "Done"],
        ["Update method statement template with rollback triggers", "Change Manager", "2026-06-15", "Done"],
        ["Review remaining POP upgrades against these lessons", "NOC Manager", "2026-10-31", "In Progress"],
    ])
    doc.save(NETOPS / "Incident_Review_Core_Router_Upgrade.docx")


def pptx_table(slide, header, rows, top=Inches(1.5)):
    shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.5), top, Inches(12.3),
                                   Inches(0.4) * (len(rows) + 1))
    t = shape.table
    for r, values in enumerate([header, *rows]):
        for c, value in enumerate(values):
            cell = t.cell(r, c)
            cell.text = str(value)
            font = cell.text_frame.paragraphs[0].runs[0].font
            font.size = PptPt(14)
            if r == 0:
                font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = PptRGB.from_string(NAVY)


def cab_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = "Halcyon Group (demo)"
    prs.core_properties.created = FIXED_TS
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Change Advisory Board - Week 43"
    s.placeholders[1].text = "Halcyon Group Network Operations | 19-25 October 2026"

    def titled(text):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = text
        return slide

    pptx_table(titled("Changes This Week"), ["Change", "Summary", "Risk", "Start", "Status"],
               [(c[0], c[1], c[2], c[3], c[6]) for c in CHANGES], top=Inches(1.4))
    high = [c for c in CHANGES if c[2] == "High"]
    pptx_table(titled("High-Risk Changes"), ["Change", "Summary", "Site", "Owner", "Status"],
               [(c[0], c[1], c[4], c[5], c[6]) for c in high])
    pptx_table(titled("Decisions and Actions"), ["Item", "Decision / Action", "Owner"], [
        ("CHG-1042", "Approval pending: confirm staged config diff and validation script results", "Network Eng Lead"),
        ("CHG-1042", "Review lessons from INC-2207 (Melbourne POP, May 2026)", "Change Manager"),
        ("CHG-1046", "Approval pending: confirm peer readiness with Sydney IX", "Peering Manager"),
        ("All", "Low and medium risk changes approved as scheduled", "CAB Chair"),
    ])
    prs.save(NETOPS / "CAB_Review_Week_43.pptx")


def main():
    for d in (LEGAL, FINANCE, NETOPS, LISTS):
        d.mkdir(parents=True, exist_ok=True)
    nda_template()
    msa(3)
    msa(4)
    close_tracker()
    board_report()
    change_register()
    method_statement()
    incident_review()
    cab_deck()
    zip_path = OUT / "Halcyon_Demo_Files.zip"
    zip_path.unlink(missing_ok=True)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        files = sorted(f for f in OUT.rglob("*") if f.is_file() and f != zip_path)
        for f in files:
            z.write(f, f.relative_to(OUT))
    print(f"Wrote {len(files)} files and {zip_path.name} to {OUT}")


if __name__ == "__main__":
    main()
