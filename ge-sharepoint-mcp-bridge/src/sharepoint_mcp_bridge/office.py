# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Richard He
"""Build, read and edit Office/text files in memory, preserving formatting."""

import csv
import difflib
import io
import json
import re
import zipfile
from pathlib import PurePosixPath
from typing import Any

import defusedxml

defusedxml.defuse_stdlib()

from docx import Document  # noqa: E402  (after defuse_stdlib on purpose)
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402
from pptx import Presentation  # noqa: E402
from pypdf import PdfReader  # noqa: E402

MAX_BYTES = 10 * 1024 * 1024
MAX_UNZIPPED = 100 * 1024 * 1024
OOXML = {".docx", ".xlsx", ".pptx"}
TEXT = {".md", ".txt", ".csv", ".json"}
CREATABLE = OOXML | TEXT
_BAD_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
PLACEHOLDER = re.compile(r"\{\{\s*(.+?)\s*\}\}")


# ---------- validation ----------

def sanitize_name(name: str) -> str:
    """A single file/folder name: no separators, traversal or reserved characters."""
    name = (name or "").strip()
    if not name or name in {".", ".."} or ".." in name or _BAD_NAME.search(name) or len(name) > 200:
        raise ValueError(f"Invalid name: {name!r}")
    return name


def sanitize_folder(path: str | None) -> str:
    """Relative folder path like 'Reports/2026'. Empty means drive root."""
    parts = [p for p in (path or "").replace("\\", "/").strip().strip("/").split("/") if p.strip()]
    return "/".join(sanitize_name(p) for p in parts)


def ext(name: str) -> str:
    return PurePosixPath(name).suffix.lower()


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
            raise ValueError("Document expands to more than 100 MB; refusing to open.")


# ---------- text extraction ----------

def _xlsx_text(data: bytes) -> list[str]:
    # openpyxl-written files have no cached formula results: show the formula instead.
    values, formulas = load_workbook(io.BytesIO(data), data_only=True), load_workbook(io.BytesIO(data))
    out = []
    for ws in values.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for vrow, frow in zip(ws.iter_rows(values_only=True), formulas[ws.title].iter_rows(values_only=True)):
            cells = [v if v is not None else f for v, f in zip(vrow, frow)]
            if any(c is not None for c in cells):
                out.append(" | ".join("" if c is None else str(c) for c in cells))
    return out


def _docx_text(data: bytes) -> list[str]:
    out = []
    for block in Document(io.BytesIO(data)).iter_inner_content():
        if hasattr(block, "rows"):
            out += [" | ".join(c.text.strip() for c in row.cells) for row in block.rows]
            continue
        text, style = block.text.strip(), (block.style.name if block.style is not None else "")
        if not text:
            continue
        level = style.removeprefix("Heading ")
        out.append(f"{'#' * int(level)} {text}" if level.isdigit() else f"# {text}" if style == "Title" else text)
    return out


def _pptx_text(data: bytes) -> list[str]:
    out = []
    for n, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        out.append(f"## Slide {n}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                out += [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
    return out


def extract_text(name: str, data: bytes, max_chars: int = 100_000) -> str:
    e = ext(name)
    if len(data) > MAX_BYTES:
        raise ValueError("File is larger than 10 MB.")
    if e in OOXML:
        _check_zip(data)
        lines = {".docx": _docx_text, ".xlsx": _xlsx_text, ".pptx": _pptx_text}[e](data)
    elif e == ".pdf":
        lines = [f"## Page {n}\n{p.extract_text() or ''}" for n, p in enumerate(PdfReader(io.BytesIO(data)).pages, 1)]
    elif e in TEXT or e in {".html", ".xml", ".log", ".yaml", ".yml"}:
        lines = [data.decode("utf-8", errors="replace")]
    else:
        raise ValueError(f"Cannot extract text from '{e}' files.")
    text = "\n".join(lines)
    return text if len(text) <= max_chars else text[:max_chars] + f"\n...[truncated at {max_chars} chars]"


# ---------- run-preserving replace ----------

def pairs(replacements: Any = None, find_text: str | None = None, replace_text: str | None = None) -> list[tuple[str, str]]:
    """Accept {"old": "new"}, [{"find","replace"}] / [{"key","value"}], or find_text/replace_text."""
    out = []
    if isinstance(replacements, dict):
        out += [(str(k), str(v)) for k, v in replacements.items() if k and v is not None]
    for d in replacements if isinstance(replacements, list) else []:
        if isinstance(d, dict):
            old = d.get("find") or d.get("key") or d.get("old") or d.get("search")
            new = next((d[k] for k in ("replace", "value", "new") if d.get(k) is not None), None)
            if old and new is not None:
                out.append((str(old), str(new)))
    if find_text and replace_text is not None:
        out.append((find_text, replace_text))
    return out


def flex_regex(text: str) -> re.Pattern | None:
    """Match 'EUR 50,000,000' from LLM-mangled keys like '_EUR_50_000_000_' (case-insensitive)."""
    tokens = re.findall(r"[^\W_]+", text)
    if not tokens:
        return None
    return re.compile(r"(?<![^\W_])" + r"[\s,._\-]+".join(map(re.escape, tokens)) + r"(?![^\W_])", re.I)


def _sub_runs(runs: list, rx: re.Pattern, repl: str) -> int:
    """Replace regex matches across runs; replacement inherits the first matched run's formatting."""
    matches = list(rx.finditer("".join(r.text for r in runs)))
    for m in reversed(matches):
        start, end, pos, first = m.start(), m.end(), 0, True
        for r in runs:
            t = r.text
            rs, re_, pos = pos, pos + len(t), pos + len(t)
            if re_ <= start or rs >= end:
                continue
            r.text = t[: max(start, rs) - rs] + (repl if first else "") + t[min(end, re_) - rs:]
            first = False
    return len(matches)


def _replace(paragraphs, find: str, repl: str) -> int:
    paragraphs = [p for p in paragraphs if p.runs]
    exact = re.compile(re.escape(find))
    n = sum(_sub_runs(p.runs, exact, repl) for p in paragraphs)
    if n == 0 and (rx := flex_regex(find)):
        n = sum(_sub_runs(p.runs, rx, repl) for p in paragraphs)
    return n


def _docx_paragraphs(doc) -> list:
    def walk(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    out = list(walk(doc))
    for s in doc.sections:
        for part in (s.header, s.footer):
            if not part.is_linked_to_previous:
                out += list(walk(part))
    return out


def _pptx_paragraphs(prs) -> list:
    def shapes(coll):
        for sh in coll:
            if sh.shape_type == 6:  # group
                yield from shapes(sh.shapes)
            else:
                yield sh

    out = []
    for slide in prs.slides:
        for sh in shapes(slide.shapes):
            if sh.has_text_frame:
                out += sh.text_frame.paragraphs
            if getattr(sh, "has_table", False):
                out += [p for row in sh.table.rows for c in row.cells for p in c.text_frame.paragraphs]
    return out


def _save(obj) -> bytes:
    buf = io.BytesIO()
    obj.save(buf)
    return buf.getvalue()


def _add_slide(prs, title: str, bullets: list | None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1 if len(prs.slide_layouts) > 1 else 0])
    if slide.shapes.title is not None:
        slide.shapes.title.text = str(title or "")
    body = [p for p in slide.placeholders if p.placeholder_format.idx != 0]
    if bullets and body:
        tf = body[0].text_frame
        tf.text = str(bullets[0])
        for b in bullets[1:]:
            tf.add_paragraph().text = str(b)


def _fill_table(table, rows) -> None:
    for row in rows:
        cells = table.add_row().cells
        for cell, val in zip(cells, row):
            cell.text = "" if val is None else str(val)


def edit_docx(data: bytes, replacements: list[tuple[str, str]], append_heading: str | None = None,
              append_paragraphs: list[str] | None = None, table_rows: list[list] | None = None,
              table_index: int = 0) -> tuple[bytes, dict]:
    _check_zip(data)
    doc = Document(io.BytesIO(data))
    paras = _docx_paragraphs(doc)
    counts = {old: _replace(paras, old, new) for old, new in replacements}
    if append_heading:
        doc.add_heading(append_heading, level=2)
    for text in append_paragraphs or []:
        doc.add_paragraph(str(text))
    if table_rows:
        if not 0 <= table_index < len(doc.tables):
            raise ValueError(f"Document has {len(doc.tables)} table(s); table_index {table_index} is out of range.")
        _fill_table(doc.tables[table_index], table_rows)
    return _save(doc), {"replacements": counts, "not_found": [k for k, v in counts.items() if not v],
                        "paragraphs_appended": len(append_paragraphs or []), "table_rows_appended": len(table_rows or [])}


def edit_pptx(data: bytes, replacements: list[tuple[str, str]], add_slides: list[dict] | None = None) -> tuple[bytes, dict]:
    _check_zip(data)
    prs = Presentation(io.BytesIO(data))
    paras = _pptx_paragraphs(prs)
    counts = {old: _replace(paras, old, new) for old, new in replacements}
    for s in add_slides or []:
        _add_slide(prs, s.get("title", ""), s.get("bullets"))
    return _save(prs), {"replacements": counts, "not_found": [k for k, v in counts.items() if not v],
                        "slides_added": len(add_slides or []), "total_slides": len(prs.slides)}


def edit_xlsx(data: bytes | None, sheet: str | None = None, initial_rows: list[list] | None = None,
              cell_updates: dict[str, Any] | None = None, append_rows: list[list] | None = None,
              add_sheet: str | None = None, add_column: dict | None = None,
              create_table: dict | None = None) -> tuple[bytes, dict]:
    """All workbook changes in one pass. data=None starts a new workbook."""
    if data:
        _check_zip(data)
    wb = load_workbook(io.BytesIO(data)) if data else Workbook()
    if not data and sheet:
        wb.active.title = sheet[:31]
    if add_sheet and add_sheet[:31] not in wb.sheetnames:
        wb.create_sheet(add_sheet[:31])
    if sheet:
        ws = wb[sheet[:31]] if sheet[:31] in wb.sheetnames else wb.create_sheet(sheet[:31])
    else:
        ws = wb[add_sheet[:31]] if add_sheet else wb.active
    empty = ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None
    if initial_rows and empty:  # explicit cells: ws.append would start at row 2 after the A1 probe
        for r, row in enumerate(initial_rows, 1):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
    for ref, val in (cell_updates or {}).items():
        ref = ref.strip().upper()
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", ref):
            raise ValueError(f"Invalid cell reference: {ref!r}")
        ws[ref] = val
    for row in append_rows or []:
        ws.append(list(row))
    if add_column:
        col = 1 if ws.max_column == 1 and ws["A1"].value is None else ws.max_column + 1
        ws.cell(row=1, column=col, value=add_column.get("header"))
        for i, val in enumerate(add_column.get("values") or [], 2):
            ws.cell(row=i, column=col, value=val)
    table = None
    if create_table:
        name = re.sub(r"\W", "_", str(create_table.get("name") or "Table1"))
        name = name if name[0].isalpha() else f"T_{name}"
        ref = create_table.get("range") or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*:[A-Z]{1,3}[1-9]\d*", ref.upper()):
            raise ValueError(f"Invalid table range: {ref!r}")
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        table = {"name": name, "range": ref}
    return _save(wb), {"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                       "cells_updated": len(cell_updates or {}), "rows_appended": len(append_rows or []),
                       "table_created": table}


# ---------- create / template ----------

def build_file(name: str, content: str = "", title: str | None = None, rows: list[list] | None = None,
               slides: list[dict] | None = None) -> bytes:
    e = ext(name)
    if e not in CREATABLE:
        raise ValueError(f"Can create {sorted(CREATABLE)} files, not '{e}'.")
    if e == ".docx":
        doc = Document()
        if title:
            doc.add_heading(title, level=0)
        for line in content.splitlines():
            s = line.strip()
            if m := re.match(r"^(#{1,3})\s+(.*)", s):
                doc.add_heading(m[2], level=len(m[1]))
            elif s.startswith(("- ", "* ")):
                doc.add_paragraph(s[2:], style="List Bullet")
            elif s:
                doc.add_paragraph(s)
        if rows:
            table = doc.add_table(rows=0, cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            _fill_table(table, rows)
        return _save(doc)
    if e == ".xlsx":
        if not rows and content.strip():
            rows = list(csv.reader(io.StringIO(content.strip())))
        return edit_xlsx(None, sheet=title or "Sheet1", initial_rows=rows)[0]
    if e == ".pptx":
        prs = Presentation()
        first = prs.slides.add_slide(prs.slide_layouts[0])
        first.shapes.title.text = title or PurePosixPath(name).stem
        if not slides and content.strip():
            slides = [{"title": title or "", "bullets": [ln.strip() for ln in content.splitlines() if ln.strip()]}]
        for s in slides or []:
            _add_slide(prs, s.get("title", ""), s.get("bullets"))
        return _save(prs)
    if e == ".csv" and rows:
        buf = io.StringIO()
        csv.writer(buf).writerows(rows)
        return buf.getvalue().encode()
    if e == ".json" and content.strip():
        json.loads(content)  # reject invalid JSON early
    return content.encode()


def fill_template(name: str, data: bytes, values: dict[str, str]) -> tuple[bytes, list[str], list[str]]:
    """Replace {{key}} placeholders (spacing/case/mangling tolerant). Returns (bytes, filled, unfilled)."""
    e = ext(name)
    if e not in OOXML:
        raise ValueError("Templates must be .docx, .pptx or .xlsx.")
    _check_zip(data)
    norm = lambda s: "_".join(re.findall(r"[^\W_]+", s.lower()))  # noqa: E731
    lookup = {norm(k): str(v) for k, v in values.items()}
    filled: set[str] = set()

    def repl(m):
        key = norm(m[1])
        if key in lookup:
            filled.add(m[1])
            return lookup[key]
        return m[0]

    if e == ".xlsx":
        wb = load_workbook(io.BytesIO(data))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and "{{" in c.value:
                        c.value = PLACEHOLDER.sub(repl, c.value)
        out = _save(wb)
    else:
        doc = Document(io.BytesIO(data)) if e == ".docx" else Presentation(io.BytesIO(data))
        paras = _docx_paragraphs(doc) if e == ".docx" else _pptx_paragraphs(doc)
        for p in paras:
            text = "".join(r.text for r in p.runs)
            for m in reversed(list(PLACEHOLDER.finditer(text))):
                new = repl(m)
                if new != m[0]:
                    _sub_runs(p.runs, re.compile(re.escape(m[0])), new)
        out = _save(doc)
    unfilled = sorted(set(PLACEHOLDER.findall(extract_text(name, out, 10**7))))
    return out, sorted(filled), unfilled


def diff(a: str, b: str, a_label: str, b_label: str, max_chars: int = 20_000) -> dict:
    lines = list(difflib.unified_diff(a.splitlines(), b.splitlines(), a_label, b_label, lineterm=""))
    body = [ln for ln in lines if not ln.startswith(("---", "+++", "@@"))]
    text = "\n".join(lines)
    return {
        "added_lines": sum(ln.startswith("+") for ln in body),
        "removed_lines": sum(ln.startswith("-") for ln in body),
        "identical": not lines,
        "diff": text if len(text) <= max_chars else text[:max_chars] + "\n...[diff truncated]",
    }
